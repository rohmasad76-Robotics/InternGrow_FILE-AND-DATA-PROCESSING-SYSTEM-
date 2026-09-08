"""
File & Data Processing System
-------------------------------
A console-based Python application for file and data management.

Features:
- Read Text Files
- Write Files
- Update Files
- Delete Files
- CSV Processing (read/write)
- Excel File Reading (using OpenPyXL)
- PDF Report Generation (upgrade)
- Logging System (upgrade)
- Automatic Backup (upgrade)

Skills used: File Handling, CSV, OpenPyXL, Exception Handling, Logging
"""

import os
import csv
import shutil
import logging
from datetime import datetime

from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

BACKUP_DIR = "backups"
LOG_FILE = "activity.log"

# ---------- Logging setup ----------
logging.basicConfig(
    filename=LOG_FILE,
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
)


class FileDataProcessingSystem:
    def __init__(self):
        if not os.path.exists(BACKUP_DIR):
            os.makedirs(BACKUP_DIR)

    # ---------- Helper: backup ----------
    def backup_file(self, filepath):
        """Copies a file into the backups/ folder before modifying/deleting it."""
        if os.path.exists(filepath):
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = os.path.basename(filepath)
            backup_path = os.path.join(BACKUP_DIR, f"{timestamp}_{filename}")
            try:
                shutil.copy2(filepath, backup_path)
                logging.info(f"Backup created for '{filepath}' at '{backup_path}'")
            except IOError as e:
                logging.error(f"Backup failed for '{filepath}': {e}")

    # ---------- Text file operations ----------
    def read_text_file(self):
        filepath = input("Enter path of the text file to read: ").strip()
        try:
            with open(filepath, "r") as f:
                content = f.read()
            print("\n--- File Content ---")
            print(content if content else "(File is empty)")
            print("--- End of File ---\n")
            logging.info(f"Read file '{filepath}'")
        except FileNotFoundError:
            print("Error: File not found.")
            logging.error(f"Read failed - file not found: '{filepath}'")
        except IOError as e:
            print(f"Error reading file: {e}")
            logging.error(f"Read failed for '{filepath}': {e}")

    def write_text_file(self):
        filepath = input("Enter path of the file to write (new or overwrite): ").strip()
        if os.path.exists(filepath):
            self.backup_file(filepath)
        print("Enter content (type 'END' on a new line to finish):")
        lines = []
        while True:
            line = input()
            if line.strip().upper() == "END":
                break
            lines.append(line)
        try:
            with open(filepath, "w") as f:
                f.write("\n".join(lines))
            print(f"File '{filepath}' written successfully.")
            logging.info(f"Wrote file '{filepath}'")
        except IOError as e:
            print(f"Error writing file: {e}")
            logging.error(f"Write failed for '{filepath}': {e}")

    def update_text_file(self):
        filepath = input("Enter path of the file to update (append to): ").strip()
        if not os.path.exists(filepath):
            print("Error: File does not exist. Use 'Write File' to create it first.")
            return
        self.backup_file(filepath)
        print("Enter content to append (type 'END' on a new line to finish):")
        lines = []
        while True:
            line = input()
            if line.strip().upper() == "END":
                break
            lines.append(line)
        try:
            with open(filepath, "a") as f:
                f.write("\n" + "\n".join(lines))
            print(f"File '{filepath}' updated successfully.")
            logging.info(f"Updated file '{filepath}'")
        except IOError as e:
            print(f"Error updating file: {e}")
            logging.error(f"Update failed for '{filepath}': {e}")

    def delete_file(self):
        filepath = input("Enter path of the file to delete: ").strip()
        if not os.path.exists(filepath):
            print("Error: File does not exist.")
            return
        confirm = input(f"Are you sure you want to delete '{filepath}'? (y/n): ").strip().lower()
        if confirm == "y":
            self.backup_file(filepath)
            try:
                os.remove(filepath)
                print(f"File '{filepath}' deleted successfully (backup saved).")
                logging.info(f"Deleted file '{filepath}'")
            except IOError as e:
                print(f"Error deleting file: {e}")
                logging.error(f"Delete failed for '{filepath}': {e}")
        else:
            print("Deletion cancelled.")

    # ---------- CSV processing ----------
    def read_csv(self):
        filepath = input("Enter path of the CSV file to read: ").strip()
        try:
            with open(filepath, "r", newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
            if not rows:
                print("CSV file is empty.")
                return
            for row in rows:
                print(row)
            logging.info(f"Read CSV file '{filepath}'")
        except FileNotFoundError:
            print("Error: File not found.")
            logging.error(f"CSV read failed - file not found: '{filepath}'")
        except csv.Error as e:
            print(f"Error reading CSV: {e}")
            logging.error(f"CSV read failed for '{filepath}': {e}")

    def write_csv(self):
        filepath = input("Enter path of the CSV file to write: ").strip()
        if os.path.exists(filepath):
            self.backup_file(filepath)
        try:
            num_cols = int(input("How many columns? "))
            headers = [input(f"Header for column {i + 1}: ").strip() for i in range(num_cols)]
        except ValueError:
            print("Invalid number entered.")
            return

        rows = [headers]
        print("Enter rows of data. Type 'done' as the first value to stop.")
        while True:
            row = []
            first = input(f"{headers[0]}: ").strip()
            if first.lower() == "done":
                break
            row.append(first)
            for h in headers[1:]:
                row.append(input(f"{h}: ").strip())
            rows.append(row)

        try:
            with open(filepath, "w", newline="") as f:
                writer = csv.writer(f)
                writer.writerows(rows)
            print(f"CSV file '{filepath}' written successfully.")
            logging.info(f"Wrote CSV file '{filepath}'")
        except IOError as e:
            print(f"Error writing CSV: {e}")
            logging.error(f"CSV write failed for '{filepath}': {e}")

    # ---------- Excel reading ----------
    def read_excel(self):
        filepath = input("Enter path of the Excel file to read: ").strip()
        try:
            wb = load_workbook(filepath, data_only=True)
            sheet = wb.active
            print(f"\n--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                print(row)
            print("--- End of Sheet ---\n")
            logging.info(f"Read Excel file '{filepath}'")
        except FileNotFoundError:
            print("Error: File not found.")
            logging.error(f"Excel read failed - file not found: '{filepath}'")
        except Exception as e:
            print(f"Error reading Excel file: {e}")
            logging.error(f"Excel read failed for '{filepath}': {e}")

    # ---------- PDF report generation ----------
    def generate_pdf_report(self):
        csv_path = input("Enter path of the CSV file to summarize: ").strip()
        pdf_path = input("Enter output PDF filename (e.g., report.pdf): ").strip()

        try:
            with open(csv_path, "r", newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
        except FileNotFoundError:
            print("Error: CSV file not found.")
            logging.error(f"PDF report failed - CSV not found: '{csv_path}'")
            return

        if not rows:
            print("CSV file is empty. Nothing to report.")
            return

        try:
            c = canvas.Canvas(pdf_path, pagesize=letter)
            width, height = letter
            y = height - 50

            c.setFont("Helvetica-Bold", 14)
            c.drawString(50, y, "Data Report")
            y -= 30

            c.setFont("Helvetica", 10)
            c.drawString(50, y, f"Source file: {csv_path}")
            y -= 15
            c.drawString(50, y, f"Total rows (including header): {len(rows)}")
            y -= 30

            for row in rows:
                line = " | ".join(str(cell) for cell in row)
                c.drawString(50, y, line[:100])  # truncate very long lines
                y -= 15
                if y < 50:  # start a new page if we run out of space
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = height - 50

            c.save()
            print(f"PDF report generated: '{pdf_path}'")
            logging.info(f"Generated PDF report '{pdf_path}' from '{csv_path}'")
        except IOError as e:
            print(f"Error generating PDF: {e}")
            logging.error(f"PDF generation failed: {e}")


def print_menu():
    print("\n===== FILE & DATA PROCESSING SYSTEM =====")
    print("1. Read Text File")
    print("2. Write Text File")
    print("3. Update Text File")
    print("4. Delete File")
    print("5. Read CSV File")
    print("6. Write CSV File")
    print("7. Read Excel File")
    print("8. Generate PDF Report (from CSV)")
    print("9. Exit")


def main():
    system = FileDataProcessingSystem()
    logging.info("Application started.")

    while True:
        print_menu()
        choice = input("Enter your choice (1-9): ").strip()

        try:
            if choice == "1":
                system.read_text_file()
            elif choice == "2":
                system.write_text_file()
            elif choice == "3":
                system.update_text_file()
            elif choice == "4":
                system.delete_file()
            elif choice == "5":
                system.read_csv()
            elif choice == "6":
                system.write_csv()
            elif choice == "7":
                system.read_excel()
            elif choice == "8":
                system.generate_pdf_report()
            elif choice == "9":
                print("Exiting File & Data Processing System. Goodbye!")
                logging.info("Application closed.")
                break
            else:
                print("Invalid choice. Please enter a number from 1 to 9.")
        except Exception as e:
            print(f"An unexpected error occurred: {e}")
            logging.error(f"Unexpected error: {e}")


if __name__ == "__main__":
    main()